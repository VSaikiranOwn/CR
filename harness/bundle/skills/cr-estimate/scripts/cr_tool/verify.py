"""Post-generation checks on the produced workbook."""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from . import constants as K

ERROR_TOKENS = ("#REF!", "#NAME?", "#VALUE!", "#DIV/0!", "#N/A", "#NULL!", "#NUM!")


def find_formula_errors(path: str | Path) -> list[str]:
    """Locate any cell whose literal content is an Excel error token."""
    wb = load_workbook(Path(path))
    found: list[str] = []
    for name in (K.SHEET_TSP, K.SHEET_DETAILS, K.SHEET_SUMMARY):
        ws = wb[name]
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                text = getattr(value, "text", value)
                if isinstance(text, str) and any(token in text for token in ERROR_TOKENS):
                    found.append(f"{name}!{cell.coordinate}")
    return found


def find_placeholders(path: str | Path) -> list[str]:
    """List cells still carrying a [TBD: ...] / [Assumption: ...] marker."""
    wb = load_workbook(Path(path))
    ws = wb[K.SHEET_DETAILS]
    found: list[str] = []
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and (
                "[TBD" in cell.value or "[Assumption" in cell.value
            ):
                found.append(f"{K.SHEET_DETAILS}!{cell.coordinate}")
    return found
