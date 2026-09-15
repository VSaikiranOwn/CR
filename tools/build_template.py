#!/usr/bin/env python3
"""Derive the blank CR framework template from a filled ver_2 workbook.

One-time (and re-runnable) build step. It takes a real, filled
``*_CR_Estimation_Framework_*.xlsx`` and strips every CR-specific value out of
the ``Tender Story Points``, ``Details`` and ``Summary`` sheets while keeping
the styling, merged ranges, column widths, dropdown validations, frozen panes
and the static ``Readme (Definitions)`` / ``MasterData`` sheets untouched.

It also appends the two Summary ratio constants to ``MasterData`` so the
Summary sheet can reference them instead of hard-coding percentages.

Usage:
    python tools/build_template.py <filled.xlsx> [templates/cr_framework_template.xlsx]
"""
from __future__ import annotations

import sys
from pathlib import Path

from copy import copy

from openpyxl import load_workbook

# Rows that carry sample data in a filled workbook. Values are cleared, styles kept.
DETAILS_DATA_ROWS = range(3, 40)
TSP_DATA_ROWS = range(3, 40)
SUMMARY_VALUE_CELLS = ["C3", "C4", "C5", "C6", "D3", "D4", "D5", "D6"]

# Appended to MasterData so Summary QA / R&D stay data-driven and self-documenting.
MASTERDATA_EXTRA = [
    (11, "QA (% of Development)", 0.4, "*Summary sheet"),
    (12, "Req. & Design (% of Development)", 0.2, "*Summary sheet"),
]


def clear_values(ws, rows, max_col):
    """Blank out cell values but leave the cell (and therefore its style) in place."""
    for r in rows:
        for c in range(1, max_col + 1):
            ws.cell(row=r, column=c).value = None


def main() -> int:
    if len(sys.argv) < 2:
        print(__doc__)
        return 1

    source = Path(sys.argv[1])
    dest = Path(sys.argv[2]) if len(sys.argv) > 2 else Path("templates/cr_framework_template.xlsx")

    wb = load_workbook(source)

    missing = [s for s in ("Readme (Definitions)", "MasterData", "Tender Story Points",
                           "Details", "Summary") if s not in wb.sheetnames]
    if missing:
        print(f"ERROR: source workbook is missing sheet(s): {', '.join(missing)}")
        return 2

    clear_values(wb["Details"], DETAILS_DATA_ROWS, 18)
    clear_values(wb["Tender Story Points"], TSP_DATA_ROWS, 26)

    summary = wb["Summary"]
    for coord in SUMMARY_VALUE_CELLS:
        summary[coord].value = None

    md = wb["MasterData"]
    for row, name, value, note in MASTERDATA_EXTRA:
        md.cell(row=row, column=1, value=name).font = copy(md["A9"].font)
        md.cell(row=row, column=2, value=value).font = copy(md["B9"].font)
        md.cell(row=row, column=2).number_format = "0%"
        if note:
            md.cell(row=row, column=3, value=note).font = copy(md["C10"].font)

    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    print(f"Template written to {dest}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
