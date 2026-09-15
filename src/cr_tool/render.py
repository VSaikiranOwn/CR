"""Fill the blank framework template from a validated spec.

The template is never rebuilt from scratch: it is loaded, its own styles are
reused, and only values and formulas are written. Everything the workbook
computes stays a live Excel formula so the sheet remains auditable and editable
after delivery.
"""
from __future__ import annotations

import math
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

from . import constants as K
from .spec import Spec

TEMPLATE_PATH = Path(__file__).resolve().parent.parent.parent / "templates" / K.TEMPLATE_NAME


# ---------------------------------------------------------------------------
# style plumbing
# ---------------------------------------------------------------------------

def _capture_row_style(ws, row: int, last_col: int) -> dict[int, dict]:
    """Snapshot the styling of a donor row so it can be replayed onto new rows."""
    captured = {}
    for col in range(1, last_col + 1):
        cell = ws.cell(row=row, column=col)
        captured[col] = {
            "font": copy(cell.font),
            "fill": copy(cell.fill),
            "border": copy(cell.border),
            "alignment": copy(cell.alignment),
            "number_format": cell.number_format,
        }
    return captured


def _apply_style(cell, style: dict) -> None:
    cell.font = copy(style["font"])
    cell.fill = copy(style["fill"])
    cell.border = copy(style["border"])
    cell.alignment = copy(style["alignment"])
    cell.number_format = style["number_format"]


def _blank_cell(cell) -> None:
    """Strip a leftover template cell back to nothing, borders included."""
    cell.value = None
    cell.font = Font()
    cell.fill = PatternFill()
    cell.border = Border()
    cell.alignment = Alignment()
    cell.number_format = "General"


def _clear_trailing(ws, first_row: int, last_col: int, keep_through: int) -> None:
    for row in range(keep_through + 1, ws.max_row + 1):
        if row < first_row:
            continue
        for col in range(1, last_col + 1):
            _blank_cell(ws.cell(row=row, column=col))
        if row in ws.row_dimensions:
            ws.row_dimensions[row].height = None


# ---------------------------------------------------------------------------
# row heights
# ---------------------------------------------------------------------------

def _wrapped_lines(text: str, width: float) -> int:
    if not text:
        return 1
    chars_per_line = max(int(width * 1.05), 8)
    lines = 0
    for segment in str(text).split("\n"):
        lines += max(1, math.ceil(len(segment) / chars_per_line))
    return lines


def _fit_height(ws, row: int, columns: list[int]) -> None:
    """Approximate Excel's auto-fit for wrapped text, which openpyxl cannot do."""
    lines = 1
    for col in columns:
        width = (ws.column_dimensions[get_column_letter(col)].width
                 or ws.sheet_format.defaultColWidth or 8.43)
        lines = max(lines, _wrapped_lines(ws.cell(row=row, column=col).value, width))
    height = min(K.MAX_ROW_HEIGHT, max(K.MIN_ROW_HEIGHT, lines * K.LINE_HEIGHT_PT + 4))
    ws.row_dimensions[row].height = height


# ---------------------------------------------------------------------------
# formulas -- these strings are the workbook's real content, keep them exact
# ---------------------------------------------------------------------------

def details_effort_formula(row: int) -> str:
    r = row
    return (
        f"=ROUND((SUM(E{r},H{r},N{r})*{K.WEIGHT_SIMPLE})"
        f"+(SUM(F{r},I{r},O{r})*{K.WEIGHT_MEDIUM})"
        f"+(SUM(G{r},J{r},P{r})*{K.WEIGHT_COMPLEX})"
        f"+(K{r}*{K.WEIGHT_DB_SIMPLE})"
        f"+(L{r}*{K.WEIGHT_DB_MEDIUM})"
        f"+(M{r}*{K.WEIGHT_DB_COMPLEX}),0)"
    )


def _size_lookup(cell_ref: str) -> str:
    return (
        f'IF({cell_ref}="Small (1-3)",1,'
        f'IF({cell_ref}="Medium (4-6)",3,'
        f'IF({cell_ref}="Large (>6)",5,0)))'
    )


def tsp_framework_unit_formula(row: int) -> str:
    r = row
    return (
        f"=({_size_lookup(f'E{r}')}*{K.MD_UI_INDEX})"
        f"+({_size_lookup(f'G{r}')}*{K.MD_PAGES_INDEX})"
        f"+({_size_lookup(f'I{r}')}*{K.MD_INTEGRATION_INDEX})"
    )


def tsp_reusability_formula(row: int) -> str:
    r = row
    choose = f"CHOOSE({{1,2,3,4,5}},L{r},N{r},P{r},R{r},T{r})"
    return (
        f'=IFERROR(SUMPRODUCT(({choose}<>"{K.NOT_APPLICABLE}")'
        f'*IFERROR(VALUE(LEFT({choose},1)),0))'
        f'/(SUMPRODUCT(({choose}<>"{K.NOT_APPLICABLE}")*1)*{K.REUSE_MAX_SCORE}),0)*100'
    )


def tsp_adjustment_formula(row: int) -> str:
    r = row
    return (
        f"=IF(V{r}<=24,{K.MD_ADJ_BAND_0},"
        f"IF(V{r}<=49,{K.MD_ADJ_BAND_25},"
        f"IF(V{r}<=74,{K.MD_ADJ_BAND_50},{K.MD_ADJ_BAND_75})))"
    )


# ---------------------------------------------------------------------------
# sheet writers
# ---------------------------------------------------------------------------

def _write_details(ws, spec: Spec) -> int:
    """Write the Details rows plus the Total row. Returns the Total row number."""
    first = K.DETAILS_FIRST_DATA_ROW
    data_style = _capture_row_style(ws, first, K.DETAILS_LAST_COL)
    total_style = _capture_row_style(ws, K.DETAILS_TEMPLATE_TOTAL_ROW, K.DETAILS_LAST_COL)

    for index, row in enumerate(spec.rows):
        r = first + index
        for col in range(1, K.DETAILS_LAST_COL + 1):
            _apply_style(ws.cell(row=r, column=col), data_style[col])

        ws.cell(row=r, column=K.DETAILS_COL["s_no"], value=row.s_no)
        ws.cell(row=r, column=K.DETAILS_COL["user_story"], value=row.user_story)
        ws.cell(row=r, column=K.DETAILS_COL["business_requirement"],
                value=row.business_requirement)
        ws.cell(row=r, column=K.DETAILS_COL["technical_component"],
                value=row.technical_component)

        for dimension in K.DETAILS_DIMENSIONS:
            base = K.DETAILS_COL[dimension]
            for offset, count in enumerate(row.counts(dimension)):
                # Blank, not 0 -- the source workbook leaves unflagged cells empty.
                ws.cell(row=r, column=base + offset, value=count or None)

        ws.cell(row=r, column=K.DETAILS_COL["dev_efforts"],
                value=details_effort_formula(r))
        _fit_height(ws, r, [K.DETAILS_COL["business_requirement"],
                            K.DETAILS_COL["technical_component"]])

    total_row = first + len(spec.rows)
    for col in range(1, K.DETAILS_LAST_COL + 1):
        _apply_style(ws.cell(row=total_row, column=col), total_style[col])
    ws.cell(row=total_row, column=K.DETAILS_COL["bpm"] + 2, value="Total")
    ws.cell(row=total_row, column=K.DETAILS_COL["dev_efforts"],
            value=f"=SUM(Q{first}:Q{total_row - 1})")
    ws.row_dimensions[total_row].height = 18

    _clear_trailing(ws, first, K.DETAILS_LAST_COL, total_row)
    return total_row


def _write_tsp(ws, spec: Spec) -> int:
    """Write one Tender Story Points row per user story. Returns the last row used."""
    first = K.TSP_FIRST_DATA_ROW
    if not spec.stories:
        _clear_trailing(ws, first, K.TSP_LAST_COL, first - 1)
        return first - 1

    data_style = _capture_row_style(ws, first, K.TSP_LAST_COL)
    detail_cols = [K.TSP_COL[f"{d}_details"]
                   for d in K.SCOPE_DIMENSIONS + K.REUSE_DIMENSIONS]

    for index, story in enumerate(spec.stories):
        r = first + index
        for col in range(1, K.TSP_LAST_COL + 1):
            _apply_style(ws.cell(row=r, column=col), data_style[col])

        ws.cell(row=r, column=K.TSP_COL["s_no"], value=index + 1)
        ws.cell(row=r, column=K.TSP_COL["epic"], value=story.epic or None)
        ws.cell(row=r, column=K.TSP_COL["us_id"], value=story.us_id)
        ws.cell(row=r, column=K.TSP_COL["us_summary"], value=story.us_summary or None)

        for dimension in K.SCOPE_DIMENSIONS:
            entry = story.scope[dimension]
            ws.cell(row=r, column=K.TSP_COL[dimension], value=entry.size)
            ws.cell(row=r, column=K.TSP_COL[f"{dimension}_details"],
                    value=entry.details or None)

        ws.cell(row=r, column=K.TSP_COL["framework_unit"],
                value=tsp_framework_unit_formula(r))

        for dimension in K.REUSE_DIMENSIONS:
            entry = story.reusability[dimension]
            ws.cell(row=r, column=K.TSP_COL[dimension], value=entry.level)
            ws.cell(row=r, column=K.TSP_COL[f"{dimension}_details"],
                    value=entry.details or None)

        score_cell = ws.cell(row=r, column=K.TSP_COL["reusability_score"])
        score_ref = f"{get_column_letter(K.TSP_COL['reusability_score'])}{r}"
        score_cell.value = ArrayFormula(score_ref, tsp_reusability_formula(r))

        ws.cell(row=r, column=K.TSP_COL["adjustment_pct"], value=tsp_adjustment_formula(r))
        ws.cell(row=r, column=K.TSP_COL["framework_unit_adj"], value=f"=K{r}*(1-W{r})")
        ws.cell(row=r, column=K.TSP_COL["req_design_unit"],
                value=f"=X{r}*({K.MD_REQ_DESIGN_UNIT})")
        ws.cell(row=r, column=K.TSP_COL["tsp"],
                value=f"=SUM(X{r},Y{r})*{K.MD_TSP_FACTOR}")

        _fit_height(ws, r, detail_cols)

    last_row = first + len(spec.stories) - 1
    _clear_trailing(ws, first, K.TSP_LAST_COL, last_row)
    _extend_validations(ws, last_row)
    return last_row


def _extend_validations(ws, last_row: int) -> None:
    """Keep the scope / reusability dropdowns covering every written row."""
    for validation in ws.data_validations.dataValidation:
        ranges = []
        for cell_range in list(validation.sqref.ranges):
            if cell_range.max_row < last_row:
                cell_range.max_row = last_row
            ranges.append(str(cell_range))
        validation.sqref = " ".join(ranges)


def _write_summary(ws, details_total_row: int, spec: Spec) -> None:
    col = K.SUMMARY_COL_TOTAL
    ws.cell(row=K.SUMMARY_ROW_DEV, column=col,
            value=f"={K.SHEET_DETAILS}!Q{details_total_row}")
    ws.cell(row=K.SUMMARY_ROW_QA, column=col,
            value=f"=ROUND(C{K.SUMMARY_ROW_DEV}*{K.MD_SUMMARY_QA_RATIO},0)")
    ws.cell(row=K.SUMMARY_ROW_RND, column=col,
            value=f"=ROUND(C{K.SUMMARY_ROW_DEV}*{K.MD_SUMMARY_RND_RATIO},0)")
    ws.cell(row=K.SUMMARY_ROW_TOTAL, column=col,
            value=f"=SUM(C{K.SUMMARY_ROW_DEV}:C{K.SUMMARY_ROW_RND})")

    original = spec.original_efforts
    ocol = K.SUMMARY_COL_ORIGINAL
    if original:
        ws.cell(row=K.SUMMARY_ROW_DEV, column=ocol, value=original.get("development"))
        ws.cell(row=K.SUMMARY_ROW_QA, column=ocol, value=original.get("qa"))
        ws.cell(row=K.SUMMARY_ROW_RND, column=ocol, value=original.get("req_design"))
        ws.cell(row=K.SUMMARY_ROW_TOTAL, column=ocol,
                value=f"=SUM(D{K.SUMMARY_ROW_DEV}:D{K.SUMMARY_ROW_RND})")
    else:
        for row in range(K.SUMMARY_ROW_DEV, K.SUMMARY_ROW_TOTAL + 1):
            ws.cell(row=row, column=ocol).value = None


# ---------------------------------------------------------------------------
# entry point
# ---------------------------------------------------------------------------

def render(spec: Spec, output: str | Path, template: str | Path | None = None) -> Path:
    """Generate the CR workbook and return the path it was written to."""
    template_path = Path(template) if template else TEMPLATE_PATH
    if not template_path.exists():
        raise FileNotFoundError(f"Template not found: {template_path}")

    wb = load_workbook(template_path)
    details_total_row = _write_details(wb[K.SHEET_DETAILS], spec)
    _write_tsp(wb[K.SHEET_TSP], spec)
    _write_summary(wb[K.SHEET_SUMMARY], details_total_row, spec)

    # Excel evaluates every formula the moment the file is opened.
    wb.calculation.fullCalcOnLoad = True

    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return output
