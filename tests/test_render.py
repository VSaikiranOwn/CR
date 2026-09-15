"""End-to-end: the generated workbook must be the ver_2 template, filled."""
import pytest
from openpyxl import load_workbook

from cr_tool import constants as K
from cr_tool.render import render
from cr_tool.spec import load
from cr_tool.verify import find_formula_errors

SAMPLE = "examples/sample_cr_spec.json"


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    out = tmp_path_factory.mktemp("cr") / "out.xlsx"
    render(load(SAMPLE), out)
    return out


@pytest.fixture(scope="module")
def wb(built):
    return load_workbook(built)


class TestWorkbookShape:
    def test_all_five_sheets_survive(self, wb):
        assert wb.sheetnames == [
            K.SHEET_README, K.SHEET_MASTERDATA, K.SHEET_TSP,
            K.SHEET_DETAILS, K.SHEET_SUMMARY,
        ]

    def test_no_formula_errors(self, built):
        assert find_formula_errors(built) == []

    def test_recalculates_on_open(self, wb):
        assert wb.calculation.fullCalcOnLoad is True

    def test_static_sheets_are_untouched(self, wb):
        assert wb[K.SHEET_README]["A1"].value == "Definitions"
        assert wb[K.SHEET_MASTERDATA]["B2"].value == 4
        assert wb[K.SHEET_MASTERDATA]["B10"].value == 1.16

    def test_summary_ratios_live_in_masterdata(self, wb):
        md = wb[K.SHEET_MASTERDATA]
        assert md["B11"].value == K.SUMMARY_QA_RATIO
        assert md["B12"].value == K.SUMMARY_RND_RATIO


class TestDetailsSheet:
    def test_row_count_and_total_row_position(self, wb):
        ws = wb[K.SHEET_DETAILS]
        # 9 sample rows starting at row 3 -> total on row 12, as in ver_2.
        assert ws["A3"].value == 1
        assert ws["P12"].value == "Total"
        assert ws["Q12"].value == "=SUM(Q3:Q11)"

    def test_effort_formula_matches_ver_2_exactly(self, wb):
        expected = ("=ROUND((SUM(E3,H3,N3)*1.5)+(SUM(F3,I3,O3)*3)"
                    "+(SUM(G3,J3,P3)*5)+(K3*0.25)+(L3*0.5)+(M3*1),0)")
        assert wb[K.SHEET_DETAILS]["Q3"].value == expected

    def test_s_no_restarts_for_the_second_story_group(self, wb):
        ws = wb[K.SHEET_DETAILS]
        assert [ws[f"A{r}"].value for r in range(3, 12)] == [1, 2, 3, 4, 5, 1, 2, 3, 4]

    def test_unflagged_cells_are_blank_not_zero(self, wb):
        ws = wb[K.SHEET_DETAILS]
        assert ws["F3"].value == 1      # UI Medium
        assert ws["E3"].value is None   # UI Simple, unflagged

    def test_rows_past_the_total_are_stripped_of_template_styling(self, wb):
        ws = wb[K.SHEET_DETAILS]
        assert ws["D13"].value is None
        left = ws["D13"].border.left
        assert left is None or left.style is None

    def test_header_styling_is_preserved(self, wb):
        header = wb[K.SHEET_DETAILS]["A1"]
        assert header.fill.fgColor.rgb == "FF0070C0"
        assert header.font.name == "Aptos Narrow"


class TestTenderStoryPointsSheet:
    def test_one_row_per_story(self, wb):
        ws = wb[K.SHEET_TSP]
        assert ws["C3"].value == "SLADCYBQSK-42132"
        assert ws["C4"].value == "SLADCYBQSK-42528"
        assert ws["C5"].value is None

    def test_dropdown_values_are_written_verbatim(self, wb):
        ws = wb[K.SHEET_TSP]
        assert ws["E3"].value in K.SCOPE_SIZES
        assert ws["L3"].value in K.REUSE_LEVELS

    def test_framework_unit_is_a_formula_over_masterdata(self, wb):
        value = wb[K.SHEET_TSP]["K3"].value
        assert value.startswith("=(IF(E3=")
        assert "MasterData!$B$2" in value
        assert "MasterData!$B$4" in value

    def test_reusability_score_is_a_single_cell_array_formula(self, wb):
        cell = wb[K.SHEET_TSP]["V3"]
        assert cell.value.ref == "V3"
        assert "SUMPRODUCT" in cell.value.text
        assert '<>"Not Applicable"' in cell.value.text

    def test_downstream_formulas(self, wb):
        ws = wb[K.SHEET_TSP]
        assert ws["X3"].value == "=K3*(1-W3)"
        assert ws["Y3"].value == "=X3*(MasterData!$B$9)"
        assert ws["Z3"].value == "=SUM(X3,Y3)*MasterData!$B$10"

    def test_dropdown_validations_survive(self, wb):
        validations = wb[K.SHEET_TSP].data_validations.dataValidation
        assert len(validations) == 2
        assert any("Small (1-3)" in v.formula1 for v in validations)
        assert any("Low Reusability" in v.formula1 for v in validations)

    def test_number_formats_are_kept(self, wb):
        ws = wb[K.SHEET_TSP]
        assert ws["W3"].number_format == "0%"
        assert ws["Z3"].number_format == "0.00"


class TestSummarySheet:
    def test_development_links_to_the_details_total(self, wb):
        assert wb[K.SHEET_SUMMARY]["C3"].value == "=Details!Q12"

    def test_qa_and_rnd_are_live_formulas(self, wb):
        ws = wb[K.SHEET_SUMMARY]
        assert ws["C4"].value == "=ROUND(C3*MasterData!$B$11,0)"
        assert ws["C5"].value == "=ROUND(C3*MasterData!$B$12,0)"
        assert ws["C6"].value == "=SUM(C3:C5)"

    def test_original_efforts_column_is_filled_when_given(self, wb):
        ws = wb[K.SHEET_SUMMARY]
        assert ws["D3"].value == 38
        assert ws["D6"].value == "=SUM(D3:D5)"

    def test_original_efforts_column_is_blank_when_omitted(self, tmp_path):
        spec = load(SAMPLE)
        spec.original_efforts = None
        out = tmp_path / "no_original.xlsx"
        render(spec, out)
        ws = load_workbook(out)[K.SHEET_SUMMARY]
        assert all(ws[f"D{r}"].value is None for r in range(3, 7))


class TestVariableRowCounts:
    @pytest.mark.parametrize("count", [1, 3, 20])
    def test_total_row_follows_the_data(self, tmp_path, count):
        spec = load(SAMPLE)
        spec.rows = (spec.rows * 20)[:count]
        for index, row in enumerate(spec.rows, 1):
            row.s_no = index
        out = tmp_path / f"n{count}.xlsx"
        render(spec, out)
        ws = load_workbook(out)[K.SHEET_DETAILS]
        total_row = 3 + count
        assert ws.cell(row=total_row, column=16).value == "Total"
        assert ws.cell(row=total_row, column=17).value == f"=SUM(Q3:Q{total_row - 1})"
        assert ws.cell(row=total_row + 1, column=17).value is None

    def test_dropdowns_extend_past_the_template_band(self, tmp_path):
        spec = load(SAMPLE)
        spec.stories = (spec.stories * 12)[:20]
        out = tmp_path / "many.xlsx"
        render(spec, out)
        ws = load_workbook(out)[K.SHEET_TSP]
        assert ws["C22"].value is not None      # 20 stories -> rows 3..22
        for validation in ws.data_validations.dataValidation:
            assert max(r.max_row for r in validation.sqref.ranges) >= 22

    def test_no_stories_leaves_the_tsp_sheet_empty_but_intact(self, tmp_path):
        spec = load(SAMPLE)
        spec.stories = []
        out = tmp_path / "nostories.xlsx"
        render(spec, out)
        ws = load_workbook(out)[K.SHEET_TSP]
        assert ws["A1"].value == "S/No"
        assert ws["C3"].value is None
